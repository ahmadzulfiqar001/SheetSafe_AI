"""XLSX algorithms for synthetic unit fixtures, never in-process user uploads."""

from contextlib import ExitStack
from datetime import date, datetime, time, timedelta
from io import BytesIO
from itertools import zip_longest
import math
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile
from zlib import error as ZlibError

from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from ingestion.contracts import IngestionCode as Code, IngestionError, IngestionRequest, MAX_COLUMNS, MAX_ROWS
from parser_worker.common import Cell, ParseControl, ProfileAccumulator, verify_source
from parser_worker.xlsx_archive import inspect_xlsx_archive


def _cell(cell, cached, profile):
    formula = cell.data_type == "f"
    value = cached.value if formula else cell.value
    if cell.data_type == "e" or cached.data_type == "e":
        profile.warnings.add("EXCEL_ERROR_VALUES")
    if value is None:
        return Cell(None, formula=formula)
    if type(value) is str:
        return Cell(value, "text", formula)
    if type(value) is bool:
        return Cell("true" if value else "false", "boolean", formula)
    if type(value) is int:
        return Cell(str(value), "integer", formula)
    if type(value) is float and math.isfinite(value):
        return Cell(repr(value), "number", formula)
    if isinstance(value, (datetime, date, time)):
        return Cell(value.isoformat(), "date", formula)
    if isinstance(value, timedelta):
        return Cell(str(value), "text", formula)
    raise IngestionError(Code.MALFORMED_XLSX)


def profile_xlsx_fixture(source: bytes, request: IngestionRequest,
        *, control: ParseControl | None = None):
    """No filesystem or DB capability. Caller supplies reviewed synthetic bytes."""
    control = control if control is not None else ParseControl()
    verify_source(source, request, control)
    if request.metadata.file_type != "xlsx" or not source.startswith(b"PK\x03\x04"):
        raise IngestionError(Code.UNSUPPORTED_CONTENT)
    try:
        layouts = inspect_xlsx_archive(source, control)
        selected = next((s for s in layouts if s.name == request.metadata.worksheet), None)
        if selected is None:
            raise IngestionError(Code.INVALID_WORKSHEET)
        if selected.rows > MAX_ROWS + 1:
            raise IngestionError(Code.ROW_LIMIT)
        if selected.columns > MAX_COLUMNS:
            raise IngestionError(Code.COLUMN_LIMIT)
        if not selected.rows or not selected.columns:
            raise IngestionError(Code.EMPTY_SOURCE)
        with ExitStack() as handles:
            workbook = load_workbook(BytesIO(source), read_only=True, data_only=False,
                keep_links=False, keep_vba=False, rich_text=False)
            handles.callback(workbook.close)
            control.check()
            cached = workbook
            if selected.formulas:
                cached = load_workbook(BytesIO(source), read_only=True, data_only=True,
                    keep_links=False, keep_vba=False, rich_text=False)
                handles.callback(cached.close)
                control.check()
            if workbook.sheetnames != [layout.name for layout in layouts]:
                raise IngestionError(Code.INVALID_WORKSHEET)
            sheet, cached_sheet = workbook[selected.name], cached[selected.name]
            sheet.reset_dimensions()
            cached_sheet.reset_dimensions()
            rows = sheet.iter_rows(min_row=1, max_row=selected.rows, max_col=selected.columns)
            handles.callback(rows.close)
            cached_rows = cached_sheet.iter_rows(min_row=1, max_row=selected.rows, max_col=selected.columns) if selected.formulas else rows
            if selected.formulas:
                handles.callback(cached_rows.close)
            headers = next(rows)
            if any(c.data_type == "f" for c in headers):
                raise IngestionError(Code.INVALID_HEADER)
            profile = ProfileAccumulator(tuple(c.value for c in headers), control)
            if selected.formulas:
                next(cached_rows)
                pairs = zip_longest(rows, cached_rows)
            else:
                pairs = ((row, row) for row in rows)
            for row, cached_row in pairs:
                if row is None or cached_row is None or len(row) != len(cached_row):
                    raise IngestionError(Code.MALFORMED_XLSX)
                profile.add(tuple(_cell(c, v, profile) for c, v in zip(row, cached_row)))
            return profile.finish(request, delimiter=None)
    except IngestionError:
        raise
    except InterruptedError:
        raise IngestionError(Code.INTERRUPTED) from None
    except DefusedXmlException:
        raise IngestionError(Code.UNSAFE_ARCHIVE) from None
    except (BadZipFile, ZlibError, ParseError, InvalidFileException, ValueError, TypeError, KeyError, IndexError, OSError, OverflowError):
        raise IngestionError(Code.MALFORMED_XLSX) from None
